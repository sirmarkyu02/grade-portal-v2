import sys
import subprocess

try:
    import openpyxl
except ImportError:
    subprocess.check_call([sys.executable, '-m', 'pip', 'install', 'openpyxl'])
    import openpyxl

file_path = r'c:\Users\PC\Desktop\PROJECTS\Grade Portal v2\Grades\11 ASSH & STEM - Gen Math.xlsx'
wb = openpyxl.load_workbook(file_path, data_only=False)

print('--- Sheets ---')
for name in wb.sheetnames:
    print(name)

print('\n--- Formulas in each sheet ---')
for sheet_name in wb.sheetnames:
    sheet = wb[sheet_name]
    print(f'\nSheet: {sheet_name}')
    formulas_found = 0
    # Search the first 50 rows and 50 columns
    for row in sheet.iter_rows(min_row=1, max_row=50, min_col=1, max_col=50):
        for cell in row:
            if cell.data_type == 'f':
                print(f'{cell.coordinate}: {cell.value}')
                formulas_found += 1
            elif isinstance(cell.value, str) and str(cell.value).startswith('='):
                # Sometimes formulas aren't caught by data_type == 'f' if openpyxl parses differently
                print(f'{cell.coordinate}: {cell.value}')
                formulas_found += 1
    if formulas_found == 0:
        print('No formulas found in the first 50x50 area.')
